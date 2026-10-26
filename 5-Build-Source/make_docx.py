#!/usr/bin/env python3
"""Converts site/_docx.html to a Word file with pandoc, then applies the brand typography and colours."""
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "deliverables" / "DAF-Act-Tokeniser-Explainer.docx"
NAVY = RGBColor(0x28, 0x3A, 0x46)
RED = RGBColor(0xFF, 0x50, 0x50)
FONT = "Readex Pro"

subprocess.run(["pandoc", str(ROOT / "site" / "_docx.html"), "-f", "html", "-t", "docx", "-o", str(OUT),
                "--metadata", "title=Australia's Digital Assets Framework Act: What It Means for Tokeniser"], check=True)

doc = Document(str(OUT))
core = doc.core_properties
core.author = ""
core.last_modified_by = ""
core.comments = "General information only, not legal advice. As at 26 October 2026."


def set_font(style, size=None, color=None, bold=None):
    style.font.name = FONT
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), FONT)
    if size:
        style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold


styles = {s.name: s for s in doc.styles}
for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
    if name in styles:
        set_font(styles[name], 10.5, NAVY)
for name, size, color in (("Title", 24, NAVY), ("Heading 1", 20, NAVY), ("Heading 2", 15, NAVY), ("Heading 3", 12, RED)):
    if name in styles:
        set_font(styles[name], size, color, bold=False)
if "Hyperlink" in styles:
    styles["Hyperlink"].font.color.rgb = NAVY

for p in doc.paragraphs:
    for r in p.runs:
        if r.bold and r.text.startswith("[") and r.text.endswith("]"):
            r.font.color.rgb = RED

for t in doc.tables:
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = t._tbl.tblPr
    for old in tbl_pr.findall(qn("w:tblW")):
        tbl_pr.remove(old)
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), "5000")
    tblw.set(qn("w:type"), "pct")
    tbl_pr.append(tblw)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "DCDCDC")
        borders.append(el)
    tbl_pr.append(borders)
    hdr = t.rows[0]
    for cell in hdr.cells:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "283A46")
        cell._tc.get_or_add_tcPr().append(shd)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.bold = True

doc.save(str(OUT))
print("docx written:", OUT)
